#!/usr/bin/env python3
"""Generador de SIGA-LoRa_Fase1_Editable.docx con la estructura completa de 5 páginas de la guía técnica."""

import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

doc = docx.Document()

# Márgenes de 0.5 pulgadas (1.27 cm)
for section in doc.sections:
    section.top_margin = Inches(0.5)
    section.bottom_margin = Inches(0.5)
    section.left_margin = Inches(0.5)
    section.right_margin = Inches(0.5)
    section.page_width = Inches(8.5)
    section.page_height = Inches(11.0)

UNAB_BLUE = RGBColor(23, 56, 100)

def set_cell_border(cell, **kwargs):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = parse_xml(r'''
        <w:tcBorders {} >
            <w:top w:val="{}" w:sz="{}" w:space="0" w:color="{}"/>
            <w:left w:val="{}" w:sz="{}" w:space="0" w:color="{}"/>
            <w:bottom w:val="{}" w:sz="{}" w:space="0" w:color="{}"/>
            <w:right w:val="{}" w:sz="{}" w:space="0" w:color="{}"/>
        </w:tcBorders>
    '''.format(
        nsdecls('w'),
        kwargs.get('top', 'single'), kwargs.get('top_sz', '4'), kwargs.get('top_color', '000000'),
        kwargs.get('left', 'single'), kwargs.get('left_sz', '4'), kwargs.get('left_color', '000000'),
        kwargs.get('bottom', 'single'), kwargs.get('bottom_sz', '4'), kwargs.get('bottom_color', '000000'),
        kwargs.get('right', 'single'), kwargs.get('right_sz', '4'), kwargs.get('right_color', '000000')
    ))
    tcPr.append(tcBorders)

def set_cell_shading(cell, color_hex):
    shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading)

# ==========================================
# PÁGINA 1: ENCABEZADO, OBJETIVOS, DINÁMICA
# ==========================================
htable = doc.add_table(rows=1, cols=3)
htable.alignment = WD_TABLE_ALIGNMENT.CENTER
htable.autofit = False
htable.columns[0].width = Inches(1.3)
htable.columns[1].width = Inches(4.9)
htable.columns[2].width = Inches(1.3)

p_l = htable.cell(0, 0).paragraphs[0]
p_l.alignment = WD_ALIGN_PARAGRAPH.CENTER
try:
    p_l.add_run().add_picture('assets/img/isotipo-azul.png', width=Inches(1.15))
except Exception:
    p_l.add_run('[UNAB]')

p_c = htable.cell(0, 1).paragraphs[0]
p_c.alignment = WD_ALIGN_PARAGRAPH.CENTER

r = p_c.add_run('Universidad Dr. Andrés Bello\n')
r.bold = True
r.font.size = Pt(13)
r.font.color.rgb = UNAB_BLUE

r = p_c.add_run('Tecnología e Innovación\n')
r.font.size = Pt(10.5)

r = p_c.add_run('Asignatura: ')
r.bold = True
r.font.size = Pt(9.5)
r = p_c.add_run('Sistemas Digitales\n')
r.font.size = Pt(9.5)

r = p_c.add_run('GUÍA TÉCNICA DE PROPUESTAS DE PROYECTOS\n')
r.bold = True
r.font.size = Pt(10.5)
r.font.color.rgb = UNAB_BLUE

r = p_c.add_run('Duración Estimada:\n')
r.bold = True
r.font.size = Pt(8.5)

r = p_c.add_run('Fecha: ')
r.bold = True
r.font.size = Pt(8.5)
r = p_c.add_run('San Miguel, sábado 10 de octubre de 2026')
r.font.size = Pt(8.5)

p_r = htable.cell(0, 2).paragraphs[0]
p_r.alignment = WD_ALIGN_PARAGRAPH.CENTER
try:
    p_r.add_run().add_picture('assets/img/logo-facultad.png', width=Inches(1.15))
except Exception:
    p_r.add_run('[FACULTAD]')

p1 = doc.add_paragraph()
p1.paragraph_format.space_before = Pt(8)
p1.paragraph_format.space_after = Pt(2)
r = p1.add_run('OBJETIVOS DE LA ACTIVIDAD')
r.bold = True
r.font.size = Pt(11)

p1_g = doc.add_paragraph()
p1_g.paragraph_format.space_before = Pt(2)
p1_g.paragraph_format.space_after = Pt(1)
r = p1_g.add_run('Objetivo General\n')
r.bold = True
r.font.size = Pt(10)
r = p1_g.add_run('Diseñar e implementar un prototipo funcional de circuito digital integrado con microcontroladores que responda a una problemática ambiental real, aplicando simplificación booleana, compuertas lógicas, integrados combinacionales y maquetado físico.')
r.font.size = Pt(9.5)

p1_e = doc.add_paragraph()
p1_e.paragraph_format.space_before = Pt(4)
p1_e.paragraph_format.space_after = Pt(1)
r = p1_e.add_run('Objetivos Específicos')
r.bold = True
r.font.size = Pt(10)

especificos = [
    'Identificar y delimitar una problemática ambiental concreta, definiendo la población beneficiaria y el impacto del sistema propuesto.',
    'Modelar el comportamiento lógico del sistema mediante tablas de verdad, ecuaciones booleanas y la inclusión explícita de módulos multiplexores (MUX) o demultiplexores (DEMUX).',
    'Simular el circuito completo en la plataforma A+ (o Tinkercad/Logisim) previo al ensamble físico.',
    'Construir una maqueta funcional utilizando Arduino, circuitos integrados digitales (TTL/CMOS) y técnicas de ensamblaje físico (cableado, soldadura en placa perforada o PCB).',
    'Redactar un informe técnico formal bajo los estándares de publicación IEEE.'
]
for obj in especificos:
    po = doc.add_paragraph(style='List Bullet')
    po.paragraph_format.space_before = Pt(1)
    po.paragraph_format.space_after = Pt(1)
    r = po.add_run(obj)
    r.font.size = Pt(9)

p1_ind = doc.add_paragraph()
p1_ind.paragraph_format.space_before = Pt(6)
p1_ind.paragraph_format.space_after = Pt(2)
r = p1_ind.add_run('INDICACIONES GENERALES Y DINÁMICA DE TRABAJO')
r.bold = True
r.font.size = Pt(11)

indicaciones = [
    ('Conformación de Equipos: ', 'Los estudiantes trabajarán estrictamente en grupos de 4 integrantes.'),
    ('Eje Temático (Medio Ambiente): ', 'El proyecto debe resolver o monitorear un problema ambiental (gestión de residuos, ahorro energético, calidad del aire/agua, automatización de riego sustentable, detección de incendios, control de huella de carbono, entre otros).'),
    ('Validación Obligatoria de Propuesta: ', 'Antes de iniciar cualquier compra o ensamble, el grupo debe someter su propuesta formal al docente. Ningún proyecto podrá desarrollarse sin la validación y firma de aprobación previa del catedrático.'),
    ('Requisitos Técnicos Mínimos del Hardware:\n', '• Circuitería Lógica Pura: Debe incluir al menos una etapa basada en lógica combinacional discreta (compuertas lógicas integradas, tablas de verdad optimizadas por Mapas de Karnaugh).\n• Módulos MUX / DEMUX: Integración funcional de al menos un integrado multiplexor o demultiplexor (ej. 74LS151, 74LS138, 74HC4051).\n• Microcontrolador (Arduino): Procesamiento de señales de sensores ambientales y control de actuadores o despliegue de datos.\n• Maquetado y Soldadura: La entrega final requiere una maqueta física a escala donde la circuitería esté soldada de forma limpia en placas perforadas/PCB y cableada adecuadamente.')
]
for title, text in indicaciones:
    pi = doc.add_paragraph(style='List Bullet')
    pi.paragraph_format.space_before = Pt(2)
    pi.paragraph_format.space_after = Pt(2)
    r = pi.add_run(title)
    r.bold = True
    r.font.size = Pt(9)
    r = pi.add_run(text)
    r.font.size = Pt(9)

p1_fases = doc.add_paragraph()
p1_fases.paragraph_format.space_before = Pt(6)
p1_fases.paragraph_format.space_after = Pt(0)
r = p1_fases.add_run('FASES Y CRONOGRAMA DE ENTREGAS')
r.bold = True
r.font.size = Pt(10.5)

# ==========================================
# PÁGINA 2: FASE 1, INTEGRANTES, PROBLEMÁTICA
# ==========================================
doc.add_page_break()

p2_tit = doc.add_paragraph()
p2_tit.paragraph_format.space_before = Pt(0)
p2_tit.paragraph_format.space_after = Pt(2)
r = p2_tit.add_run('Fase 1: Ficha de Propuesta de Proyecto (Entrega para Validación)\n')
r.bold = True
r.font.size = Pt(11)
r = p2_tit.add_run('El equipo entregará un documento de un máximo de dos páginas que contenga:')
r.font.size = Pt(10)

p2_sec1 = doc.add_paragraph()
p2_sec1.paragraph_format.space_before = Pt(4)
p2_sec1.paragraph_format.space_after = Pt(2)
r = p2_sec1.add_run('▪  Nombre del proyecto y miembros del equipo.')
r.bold = True
r.font.size = Pt(10)

ttable = doc.add_table(rows=6, cols=4)
ttable.alignment = WD_TABLE_ALIGNMENT.CENTER
ttable.autofit = False
ttable.columns[0].width = Inches(1.8)
ttable.columns[1].width = Inches(2.7)
ttable.columns[2].width = Inches(1.0)
ttable.columns[3].width = Inches(2.0)

c0 = ttable.cell(0, 0)
c0.paragraphs[0].add_run('Nombre del Proyecto:').bold = True
c0.paragraphs[0].runs[0].font.size = Pt(9.5)
c1 = ttable.cell(0, 1)
c1.merge(ttable.cell(0, 3))
p_title = c1.paragraphs[0]
r = p_title.add_run('SIGA-LoRa: Sistema Inteligente de Gestión de Alertas\n')
r.bold = True
r.font.size = Pt(10)
r = p_title.add_run('Prototipo de monitoreo inalámbrico de niveles de agua ante riesgos de inundación.')
r.italic = True
r.font.size = Pt(8.5)

c_head = ttable.cell(1, 0)
c_head.merge(ttable.cell(1, 3))
r = c_head.paragraphs[0].add_run('Integrantes del Equipo (4 estudiantes)')
r.bold = True
r.font.size = Pt(9.5)

members_data = [
    ('Estudiante (Líder):', 'Alexis Sebastián Castro Ramírez', 'CR0367042024', False),
    ('Estudiante:', 'Jefferson Jared Bonilla Meléndez', 'BM0749042022', False),
    ('Estudiante:', 'Jaime Fernando Berrios Ortiz', 'BO0310042024', False),
    ('Estudiante:', 'Por confirmar', 'Por confirmar', True)
]

for idx, (role, name, code, is_pending) in enumerate(members_data, start=2):
    cell_r = ttable.cell(idx, 0)
    r = cell_r.paragraphs[0].add_run(role)
    r.bold = True
    r.font.size = Pt(9)
    
    cell_n = ttable.cell(idx, 1)
    r = cell_n.paragraphs[0].add_run(name)
    r.italic = is_pending
    r.font.size = Pt(9)
    
    cell_c_lbl = ttable.cell(idx, 2)
    r = cell_c_lbl.paragraphs[0].add_run('Carnet:')
    r.bold = True
    r.font.size = Pt(9)
    
    cell_code = ttable.cell(idx, 3)
    r = cell_code.paragraphs[0].add_run(code)
    r.italic = is_pending
    r.font.size = Pt(9)

for row in ttable.rows:
    for cell in row.cells:
        set_cell_border(cell)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

p2_sec2 = doc.add_paragraph()
p2_sec2.paragraph_format.space_before = Pt(6)
p2_sec2.paragraph_format.space_after = Pt(2)
r = p2_sec2.add_run('▪  Problemática ambiental identificada y justificación.\n▪  Población objetivo y beneficiarios directos.')
r.bold = True
r.font.size = Pt(10)

ptable = doc.add_table(rows=1, cols=1)
ptable.alignment = WD_TABLE_ALIGNMENT.CENTER
ptable.columns[0].width = Inches(7.5)
set_cell_border(ptable.cell(0, 0))

cell_p = ptable.cell(0, 0)
p = cell_p.paragraphs[0]
p.paragraph_format.space_before = Pt(2)
p.paragraph_format.space_after = Pt(2)

r = p.add_run('Problemática Ambiental Identificada:\n')
r.bold = True
r.font.size = Pt(9.5)

r = p.add_run(
    'En el municipio de San Miguel Centro, El Salvador, los temporales y ondas tropicales provocan dos fenómenos hidrológicos distintos: por un lado, crecidas fluviales y desbordamientos en la cuenca baja del Río Grande de San Miguel (alcanzando niveles críticos de hasta 8.52 m en la estación La Canoa, equivalentes al 89 % de banca llena [4], con alertas de Protección Civil [5]); y por otro, inundaciones urbanas repentinas derivadas de la saturación de drenajes pluviales, deficiencias de infraestructura y colapso de muros perimetrales en sectores como la colonia Palo Blanco [1], [2], [3]. Registros periodísticos y de Protección Civil documentan anegamientos en calle Los Laureles (11 viviendas tras colapso estructural en 2020 [2]) y en senda Los Laureles (más de 15 viviendas por escorrentía superficial [3]). Aunque no existe evidencia de que las inundaciones de Palo Blanco provengan directamente del río, ambos escenarios confirman la vulnerabilidad local y la carencia de monitoreo instrumental continuo para emitir advertencias oportunas.\n\n'
)
r.font.size = Pt(8.8)

r = p.add_run('Justificación:\n')
r.bold = True
r.font.size = Pt(9.5)

r = p.add_run(
    'Se propone el diseño de un prototipo académico de laboratorio para monitoreo instrumental de niveles de agua y emisión de advertencias locales. El sistema no sustituye los sistemas oficiales de protección civil ni garantiza alcances absolutos antes de pruebas de campo. Emplea un sensor ultrasónico impermeable (JSN-SR04T) y tres interruptores tipo flotador para verificar niveles de lámina de agua de forma redundante. A nivel de hardware, compuertas discretas 74HC08 (AND), 74HC32 (OR) y 74HC04 (NOT) validan físicamente las condiciones de alarma antes o en paralelo al Arduino Uno, cumpliendo el requisito de lógica pura. La telemetría entre estaciones se gestiona mediante microcontroladores ESP32 acoplados a transceptores LoRa RFM95W (915 MHz), permitiendo establecer un enlace inalámbrico punto a punto sin dependencia de red Wi-Fi o telefonía celular. En la estación receptora, un CI 74HC238 (DEMUX 3:8) decodifica el código de advertencia para conmutar selectivamente LEDs por código de color y un mini buzzer BGD08 para señalización local.\n\n'
)
r.font.size = Pt(8.8)

r = p.add_run('Población Objetivo y Beneficiarios Directos:\n')
r.bold = True
r.font.size = Pt(9.5)

r = p.add_run('Población Objetivo: ')
r.bold = True
r.font.size = Pt(8.8)
r = p.add_run('Habitantes y familias de la colonia Palo Blanco (San Miguel Centro), específicamente aquellos domicilios ubicados en sectores vulnerables a inundaciones repentinas y saturación de drenajes pluviales.\n')
r.font.size = Pt(8.8)

r = p.add_run('Beneficiarios Directos: ')
r.bold = True
r.font.size = Pt(8.8)
r = p.add_run('Cantidad de beneficiarios directos pendiente de delimitación y verificación formal (familias residentes en el área proyectada de cobertura del nodo receptor). Se aclara que la cifra de 74,650 hogares (Censo 2024 de ONEC/BCR [7]) corresponde a la totalidad distrital de San Miguel y no a Palo Blanco; asimismo, los antecedentes de 11 y más de 15 viviendas afectadas en calle y senda Los Laureles son eventos históricos independientes que sirven de referencia para el posterior levantamiento de campo.')
r.font.size = Pt(8.8)

# ==========================================
# PÁGINA 3: DIAGRAMA DE BLOQUES Y PRESUPUESTO
# ==========================================
doc.add_page_break()

p3_diag_tit = doc.add_paragraph()
p3_diag_tit.paragraph_format.space_before = Pt(0)
p3_diag_tit.paragraph_format.space_after = Pt(2)
r = p3_diag_tit.add_run('▪  Diagrama de bloques preliminar del sistema (Sensores, Lógica Combinacional / MUX / DEMUX, Arduino, Actuadores).')
r.bold = True
r.font.size = Pt(10)

p_img = doc.add_paragraph()
p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_img.paragraph_format.space_before = Pt(2)
p_img.paragraph_format.space_after = Pt(4)
try:
    p_img.add_run().add_picture('assets/img/circuito.png', width=Inches(5.8))
except Exception:
    p_img.add_run('[DIAGRAMA DE BLOQUES]')

p3_pres_tit = doc.add_paragraph()
p3_pres_tit.paragraph_format.space_before = Pt(4)
p3_pres_tit.paragraph_format.space_after = Pt(2)
r = p3_pres_tit.add_run('▪  Presupuesto estimado y lista de materiales a utilizar.')
r.bold = True
r.font.size = Pt(10)

btable = doc.add_table(rows=1, cols=6)
btable.alignment = WD_TABLE_ALIGNMENT.CENTER
btable.autofit = False
col_widths = [Inches(0.5), Inches(2.2), Inches(2.4), Inches(0.8), Inches(0.8), Inches(0.8)]
for j, w in enumerate(col_widths):
    btable.columns[j].width = w

hdr_cells = btable.rows[0].cells
headers = ['Cant.', 'Descripción del Material / Componente', 'Función / Aplicación en el Proyecto', 'P. Unit.', 'Subtotal', 'Estado']
for j, h in enumerate(headers):
    hdr_cells[j].text = h
    hdr_cells[j].paragraphs[0].runs[0].bold = True
    hdr_cells[j].paragraphs[0].runs[0].font.size = Pt(7.5)
    set_cell_shading(hdr_cells[j], 'F2F2F2')
    set_cell_border(hdr_cells[j])

budget_rows = [
    ('HEADER', 'A. Equipos y Módulos de Adquisición y Comunicaciones (Disponibles en inventario del equipo)', ''),
    ('1', 'Arduino Uno R3 (microcontrolador ATmega328P)', 'Adquisición de señales y lógica de sensores', '$5.10', '$5.10', 'Disponible'),
    ('2', 'Módulos ESP32 NodeMCU (30 pines)', 'Gestión de protocolo LoRa en nodos Tx y Rx', '$4.37', '$8.74', 'Disponibles'),
    ('2', 'HopeRF RFM95W 915 MHz (Transceptor LoRa + ant. helic.)', 'Enlace inalámbrico de telemetría a larga distancia', '$7.70', '$15.40', 'Disponibles'),
    ('2', 'Antenas LoRa omnidireccionales 6 dBi (adaptador SMA)', 'Ganancia y radiación RF para nodos Tx/Rx', '$5.33', '$10.66', 'Disponibles'),
    ('SUBTOTAL_A', 'Subtotal Componentes Disponibles (Valor referencial inventario):', '$39.90', 'Adquirido'),
    ('HEADER', 'B. Sensores, Lógica Digital Discreta, Actuadores y Montaje (Materiales proyectados por adquirir)', ''),
    ('1', 'Sensor ultrasónico impermeable JSN-SR04T', 'Medición continua de lámina de agua', '$2.12', '$2.12', 'Por adquirir*'),
    ('3', 'Interruptores de nivel tipo flotador vertical', 'Detección discreta escalonada de umbrales', '$1.29', '$3.87', 'Por adquirir*'),
    ('1', 'CI 74HC08 (4 compuertas AND 2 entradas, DIP-14)', 'Validación lógica física de nivel crítico', '$0.83', '$0.83', 'Por adquirir'),
    ('1', 'CI 74HC32 (4 compuertas OR 2 entradas, DIP-14)', 'Enrutamiento combinacional de alarmas', '$0.73', '$0.73', 'Por adquirir'),
    ('1', 'CI 74HC04 (6 inversores NOT, DIP-14; lote 5 uds)', 'Inversión de señales lógicas (1 inst., 4 res.)', '$1.92', '$1.92', 'Por adquirir'),
    ('1', 'CI 74HC238 (Decodificador/DEMUX 3:8; lote 5 uds)', 'Decodificación y selección de estados de alerta', '$1.12', '$1.12', 'Por adquirir'),
    ('4', 'Diodos LED 5 mm (1 verde, 1 amarillo, 1 naranja, 1 rojo)', 'Señalización óptica por niveles en nodo receptor', '$0.20', '$0.80', 'Por adquirir'),
    ('1', 'Mini buzzer piezoeléctrico continuo BGD08', 'Alerta acústica local ante umbral de peligro', '$1.00', '$1.00', 'Por adquirir'),
    ('2', 'Baquelas perforadas soldables mod. 400 (con pistas)', 'Base física y conexionado soldado de circuitos', '$3.49', '$6.98', 'Por adquirir'),
    ('2', 'Baquelas perforadas soldables mod. 405 (islas simples)', 'Submódulos auxiliares y desacople físico', '$1.49', '$2.98', 'Por adquirir'),
    ('1', 'Soldadura estaño-plomo SOL60-020 (tubo 17 g)', 'Unión eléctrica soldada permanente', '$3.99', '$3.99', 'Por adquirir'),
    ('8 m', 'Alambre estañado para conexiones 22 AWG', 'Cableado interno soldado entre componentes', '$0.25', '$2.00', 'Por adquirir'),
    ('Lote', 'Zócalos DIP, resistencias de polarización y capacitores', 'Protección de CI, limitación de LEDs y desacoplo', '$3.50', '$3.50', 'Por adquirir*'),
    ('2', 'Convertidores de nivel bidireccionales 3.3V / 5V', 'Acoplamiento lógico entre Arduino y ESP32/LoRa', '$1.25', '$2.50', 'Por adquirir*'),
    ('Lote', 'Accesorios de alimentación (porta baterías/USB)', 'Alimentación autónoma para estación Tx y Rx', '$6.00', '$6.00', 'Por adquirir*'),
    ('Lote', 'Materiales de maqueta y estructura de soporte', 'Simulación a escala de canal y montaje físico', '$8.00', '$8.00', 'Por adquirir*'),
    ('SUBTOTAL_B', 'Subtotal Materiales por Adquirir (Costo proyectado de compra):', '$48.34', 'Pendiente'),
    ('TOTAL', 'COSTO TOTAL ESTIMADO REFERENCIAL DEL PROYECTO:', '$88.24', '---')
]

for item in budget_rows:
    row_cells = btable.add_row().cells
    for j, w in enumerate(col_widths):
        row_cells[j].width = w
    tag = item[0]
    if tag == 'HEADER':
        row_cells[0].merge(row_cells[5])
        row_cells[0].text = item[1]
        row_cells[0].paragraphs[0].runs[0].bold = True
        row_cells[0].paragraphs[0].runs[0].font.size = Pt(7.2)
        set_cell_shading(row_cells[0], 'F9F9F9')
        set_cell_border(row_cells[0])
    elif tag.startswith('SUBTOTAL') or tag == 'TOTAL':
        row_cells[0].merge(row_cells[3])
        row_cells[0].text = item[1]
        row_cells[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
        row_cells[0].paragraphs[0].runs[0].bold = True
        row_cells[0].paragraphs[0].runs[0].font.size = Pt(7.2)
        set_cell_border(row_cells[0])
        
        row_cells[4].text = item[2]
        row_cells[4].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        row_cells[4].paragraphs[0].runs[0].bold = True
        row_cells[4].paragraphs[0].runs[0].font.size = Pt(7.2)
        set_cell_border(row_cells[4])
        
        row_cells[5].text = item[3]
        row_cells[5].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        row_cells[5].paragraphs[0].runs[0].italic = True
        row_cells[5].paragraphs[0].runs[0].font.size = Pt(7.0)
        set_cell_border(row_cells[5])
        if tag == 'TOTAL':
            set_cell_shading(row_cells[0], 'F2F2F2')
            set_cell_shading(row_cells[4], 'F2F2F2')
            set_cell_shading(row_cells[5], 'F2F2F2')
    else:
        cant, desc, func, pu, subt, est = item
        row_cells[0].text = cant
        row_cells[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        row_cells[1].text = desc
        row_cells[2].text = func
        row_cells[3].text = pu
        row_cells[3].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        row_cells[4].text = subt
        row_cells[4].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        row_cells[5].text = est
        row_cells[5].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        for c in row_cells:
            c.paragraphs[0].runs[0].font.size = Pt(7.0)
            set_cell_border(c)

p_bnote = doc.add_paragraph()
p_bnote.paragraph_format.space_before = Pt(2)
p_bnote.paragraph_format.space_after = Pt(2)
r = p_bnote.add_run('Notas de presupuesto: (*) Valores estimados referenciales pendientes de cotización local final. Se cuenta con un Arduino Nano disponible en reserva sin costo imputado. Precios de CI 74HC04 y 74HC238 corresponden a lotes mínimos comerciales LCSC (unidades sobrantes quedan como reserva técnica).')
r.italic = True
r.font.size = Pt(6.8)

# ==========================================
# PÁGINA 4: FIRMAS, FASE 2, SECCIONES I-III
# ==========================================
doc.add_page_break()

p4_sig_tit = doc.add_paragraph()
p4_sig_tit.paragraph_format.space_before = Pt(0)
p4_sig_tit.paragraph_format.space_after = Pt(2)
r = p4_sig_tit.add_run('▪  Firmas de conformidad y validación')
r.bold = True
r.font.size = Pt(10)

stable = doc.add_table(rows=1, cols=1)
stable.alignment = WD_TABLE_ALIGNMENT.CENTER
stable.columns[0].width = Inches(7.5)
set_cell_border(stable.cell(0, 0))

p_sig = stable.cell(0, 0).paragraphs[0]
p_sig.paragraph_format.space_before = Pt(2)
p_sig.paragraph_format.space_after = Pt(2)

r = p_sig.add_run('Firma del Líder del Equipo:\n')
r.bold = True
r.font.size = Pt(9.5)
r = p_sig.add_run('Fecha de entrega: _____ / _____ / _________\n\n')
r.font.size = Pt(9)

r = p_sig.add_run('RESERVADO PARA EL DOCENTE DE LA ASIGNATURA (MduEV. Ing. Franklin Geovany Buruca Hernández):\n')
r.bold = True
r.font.size = Pt(9)

r = p_sig.add_run('▪  Dictamen:  [   ] APROBADO       [   ] RECHAZADO       [   ] REVISAR Y REENVIAR\n')
r.bold = True
r.font.size = Pt(8.5)

r = p_sig.add_run('▪  Observaciones / Modificaciones sugeridas:\n\n\n')
r.bold = True
r.font.size = Pt(8.5)

r = p_sig.add_run('Firma y Sello del Catedrático\n')
r.bold = True
r.font.size = Pt(9.5)
r = p_sig.add_run('Fecha de validación: _____ / _____ / _________')
r.font.size = Pt(9)

p4_ent = doc.add_paragraph()
p4_ent.paragraph_format.space_before = Pt(6)
p4_ent.paragraph_format.space_after = Pt(2)
r = p4_ent.add_run('FECHA DE ENTREGA: SÁBADO 10 DE OCTUBRE DE 2026')
r.bold = True
r.font.size = Pt(10.5)

p4_f2 = doc.add_paragraph()
p4_f2.paragraph_format.space_before = Pt(2)
p4_f2.paragraph_format.space_after = Pt(1)
r = p4_f2.add_run('Fase 2: Informe Final y Defensa de Proyecto\n')
r.bold = True
r.font.size = Pt(10.5)
r = p4_f2.add_run('Entrega del informe escrito en formato IEEE, demostración del prototipo maquetado en vivo y presentación ejecutiva con diapositivas ante la clase.')
r.font.size = Pt(9.5)

p4_ieee = doc.add_paragraph()
p4_ieee.paragraph_format.space_before = Pt(4)
p4_ieee.paragraph_format.space_after = Pt(2)
r = p4_ieee.add_run('ESTRUCTURA DEL INFORME TÉCNICO (FORMATO IEEE)\n')
r.bold = True
r.font.size = Pt(10)
r = p4_ieee.add_run('El informe final debe redactarse respetando estrictamente la plantilla estándar IEEE a dos columnas, con una extensión de entre 6 y 10 páginas.')
r.font.size = Pt(9.5)

sec_table1 = doc.add_table(rows=1, cols=1)
sec_table1.alignment = WD_TABLE_ALIGNMENT.CENTER
sec_table1.columns[0].width = Inches(7.5)
set_cell_border(sec_table1.cell(0, 0))

p_s1 = sec_table1.cell(0, 0).paragraphs[0]
p_s1.paragraph_format.space_before = Pt(2)
p_s1.paragraph_format.space_after = Pt(2)

r = p_s1.add_run('Detalle de Secciones del Informe:\n\n')
r.bold = True
r.font.size = Pt(10)

r = p_s1.add_run('I. INTRODUCCIÓN:\n')
r.bold = True
r.font.size = Pt(9.5)
r = p_s1.add_run('Antecedentes generales de la aplicación de la lógica digital a problemas de la ingeniería.\n\n')
r.font.size = Pt(9)

r = p_s1.add_run('II. IDENTIFICACIÓN DE LA PROBLEMÁTICA Y POBLACIÓN OBJETIVO:\n')
r.bold = True
r.font.size = Pt(9.5)
r = p_s1.add_run('✓ Diagnóstico del problema ambiental (ej. desperdicio de agua en cultivos urbanos).\n✓ Justificación socioambiental.\n✓ Delimitación del segmento beneficiario.\n\n')
r.font.size = Pt(9)

r = p_s1.add_run('III. DISEÑO LÓGICO Y FUNDAMENTACIÓN TEÓRICA:\n')
r.bold = True
r.font.size = Pt(9.5)
r = p_s1.add_run('✓ Especificación formal de entradas y salidas digitales.\n✓ Tablas de verdad completas.\n✓ Ecuaciones booleanas obtenidas y su proceso de simplificación algebraico o por mapas de Karnaugh.')
r.font.size = Pt(9)

# ==========================================
# PÁGINA 5: SECCIONES IV-VI Y REFERENCIAS
# ==========================================
doc.add_page_break()

sec_table2 = doc.add_table(rows=1, cols=1)
sec_table2.alignment = WD_TABLE_ALIGNMENT.CENTER
sec_table2.columns[0].width = Inches(7.5)
set_cell_border(sec_table2.cell(0, 0))

p_s2 = sec_table2.cell(0, 0).paragraphs[0]
p_s2.paragraph_format.space_before = Pt(2)
p_s2.paragraph_format.space_after = Pt(2)

r = p_s2.add_run('✓ Implementación y diagrama esquemático con compuertas lógicas e integrados de Multiplexor/Demultiplexor.\n\n')
r.font.size = Pt(9)

r = p_s2.add_run('IV. ARQUITECTURA DEL SISTEMA Y MAQUETADO FÍSICO:\n')
r.bold = True
r.font.size = Pt(9.5)
r = p_s2.add_run('✓ Diagrama eléctrico general que incluya la etapa de procesamiento digital y la interfaz con Arduino.\n✓ Código fuente del microcontrolador (comentado y estructurado).\n✓ Proceso de construcción físico: Maquetado del entorno ambiental, técnicas de soldadura empleadas, distribución de componentes y medidas de seguridad eléctrica.\n\n')
r.font.size = Pt(9)

r = p_s2.add_run('V. SIMULACIÓN, PRUEBAS Y RESULTADOS:\n')
r.bold = True
r.font.size = Pt(9.5)
r = p_s2.add_run('✓ Resultados de la simulación en A+, Logisim o Tinkercad.\n✓ Pruebas de campo o de laboratorio del prototipo físico en condiciones reales o controladas.\n✓ Análisis comparativo entre el comportamiento simulado y el comportamiento real.\n\n')
r.font.size = Pt(9)

r = p_s2.add_run('VI. CONCLUSIONES Y TRABAJO FUTURO:\n')
r.bold = True
r.font.size = Pt(9.5)
r = p_s2.add_run('Evaluación del cumplimiento de objetivos y posibles mejoras a la plataforma tecnológica.\n\n')
r.font.size = Pt(9)

r = p_s2.add_run('REFERENCIAS: Fuentes bibliográficas en formato IEEE.\n')
r.bold = True
r.font.size = Pt(9.5)

refs_lines = (
    '[1] Dir. Gral. de Protección Civil, «Informe de afectaciones por lluvias», San Salvador, El Salvador, nov. 2014.\n'
    '[2] La Prensa Gráfica, «Lluvia inunda 11 viviendas y calle», San Miguel, El Salvador, abr. 2020.\n'
    '[3] La Prensa Gráfica, «Filtración de agua afecta residencias», San Miguel, El Salvador, sep. 2020.\n'
    '[4] MARN, «Niveles críticos en estación La Canoa, Río Grande de San Miguel», San Salvador, ago. 2021.\n'
    '[5] Dir. Gral. Protección Civil, «Alerta Verde por lluvias asociadas a Tormenta Tropical Melissa», oct. 2025.\n'
    '[6] Texas Instruments, «CD74HC238 High-Speed CMOS Logic 3-to-8 Line Decoder/Demultiplexer», Datasheet SCHS207D, 2003.\n'
    '[7] BCR / ONEC, «VII Censo de Población y VI de Vivienda 2024: Hogares por distrito», San Salvador, 2025.'
)
r = p_s2.add_run(refs_lines)
r.font.size = Pt(8.5)

doc.save('SIGA-LoRa_Fase1_Editable.docx')
print('SIGA-LoRa_Fase1_Editable.docx generado exitosamente con 5 páginas.')
