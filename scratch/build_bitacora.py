import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def create_bitacora_docx():
    doc = docx.Document()

    # Page Margins
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_inst = p_title.add_run("UNIVERSIDAD DR. ANDRÉS BELLO\n")
    r_inst.bold = True
    r_inst.font.size = Pt(13)
    r_inst.font.color.rgb = RGBColor(0x17, 0x38, 0x64)
    r_inst.font.name = 'Arial'

    r_fac = p_title.add_run("FACULTAD DE TECNOLOGÍA E INNOVACIÓN\n")
    r_fac.bold = True
    r_fac.font.size = Pt(11)
    r_fac.font.color.rgb = RGBColor(0x17, 0x38, 0x64)
    r_fac.font.name = 'Arial'

    r_sub = p_title.add_run("DISEÑO DE SISTEMAS DE SEGURIDAD DE REDES DE DATOS\n\n")
    r_sub.font.size = Pt(11)
    r_sub.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
    r_sub.font.name = 'Arial'

    r_bit = p_title.add_run("BITÁCORA DE AVANCES Y MODIFICACIONES — FASE 2")
    r_bit.bold = True
    r_bit.font.size = Pt(14)
    r_bit.font.color.rgb = RGBColor(0x17, 0x38, 0x64)
    r_bit.font.name = 'Arial'

    doc.add_paragraph()

    # Metadata Table
    meta_table = doc.add_table(rows=5, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_data = [
        ("Proyecto:", "Diseño de un modelo de seguridad perimetral y segmentación lógica para el aseguramiento de la red de datos en la Sede de Enlaces con la Educación (Santa Rosa de Lima, 2026)."),
        ("Periodo / Fase:", "Fase 2: Gestión de Riesgos y Políticas."),
        ("Fecha de entrega:", "12 de octubre de 2026."),
        ("Docente titular:", "MduEV. Ing. Franklin Geovany Buruca Hernández."),
        ("Equipo de trabajo:", "Genesis Saravia (Líder), Jaime Berrios, Jefferson Bonilla, Rosmeri González, Florentín Treminio.")
    ]

    for row_idx, (k, v) in enumerate(meta_data):
        c0, c1 = meta_table.rows[row_idx].cells
        c0.text = k
        c1.text = v
        set_cell_background(c0, "F0F4F8")
        set_cell_background(c1, "FFFFFF")
        set_cell_margins(c0, top=60, bottom=60, left=100, right=100)
        set_cell_margins(c1, top=60, bottom=60, left=100, right=100)
        c0.paragraphs[0].runs[0].bold = True
        c0.paragraphs[0].runs[0].font.name = 'Arial'
        c0.paragraphs[0].runs[0].font.size = Pt(9.5)
        c0.paragraphs[0].runs[0].font.color.rgb = RGBColor(0x17, 0x38, 0x64)
        c1.paragraphs[0].runs[0].font.name = 'Arial'
        c1.paragraphs[0].runs[0].font.size = Pt(9.5)

    doc.add_paragraph()

    # Section: Resumen
    p_sec1 = doc.add_paragraph()
    r_sec1 = p_sec1.add_run("1. Propósito de la Bitácora")
    r_sec1.bold = True
    r_sec1.font.size = Pt(11)
    r_sec1.font.color.rgb = RGBColor(0x17, 0x38, 0x64)
    r_sec1.font.name = 'Arial'

    p_desc1 = doc.add_paragraph(
        "El presente documento tiene como objetivo registrar de forma sistemática las correcciones, adaptaciones metodológicas y desarrollos técnicos ejecutados por el equipo investigador durante la Fase 2 del proyecto integrador de cátedra, atendiendo las directrices académicas y las observaciones emitidas por la cátedra en el primer entregable."
    )
    p_desc1.runs[0].font.name = 'Arial'
    p_desc1.runs[0].font.size = Pt(10)

    # Section: Tabla de Actividades y Cambios
    p_sec2 = doc.add_paragraph()
    r_sec2 = p_sec2.add_run("2. Registro de Actividades y Cambios Implementados")
    r_sec2.bold = True
    r_sec2.font.size = Pt(11)
    r_sec2.font.color.rgb = RGBColor(0x17, 0x38, 0x64)
    r_sec2.font.name = 'Arial'

    act_table = doc.add_table(rows=7, cols=3)
    act_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ["Aspecto / Módulo", "Modificación o Avance Realizado", "Impacto Metodológico"]
    rows = [
        (
            "Objetivo Específico 4",
            "Sustitución del verbo «Implementar» por «Configurar controles de capa de aplicación... en un laboratorio simulado».",
            "Aclara el alcance tecnológico sin comprometer la intervención física en las instalaciones reales de la sede."
        ),
        (
            "Carátula Institucional",
            "Ajuste de la portada para presentar un único logotipo institucional UNAB centrado y debidamente proporcionado.",
            "Elimina duplicidad visual y estandariza la presentación según el manual de estilo institucional."
        ),
        (
            "Marco Metodológico (Sampieri)",
            "Ratificación del diseño descriptivo y propositivo (no experimental y transversal) y justificación de no requerir hipótesis causal obligatoria.",
            "Alineación estricta con las pautas de Hernández-Sampieri y la retroalimentación del docente."
        ),
        (
            "Análisis de Resultados (Fase 2)",
            "Incorporación de la Matriz de Análisis y Evaluación de Riesgos formal categorizando 5 activos críticos y evaluando probabilidad e impacto.",
            "Cumplimiento formal del entregable de la Fase 2 (Gestión de Riesgos) según la planificación académica."
        ),
        (
            "Instrumentos de Medición",
            "Diseño y formalización del Cuestionario de Diagnóstico Técnico de Infraestructura y Seguridad de Red aplicado a informantes clave.",
            "Respaldo metodológico documental para la recolección de datos y justificación técnica del diagnóstico."
        ),
        (
            "Anexos y Evidencias",
            "Inclusión del registro fotográfico de las sesiones de trabajo del equipo (Anexo 1: Lectura de protocolo; Anexo 2: Marco teórico y riesgos).",
            "Constancia y verificación del trabajo colaborativo de los integrantes del equipo."
        )
    ]

    for col_idx, h in enumerate(headers):
        c = act_table.rows[0].cells[col_idx]
        c.text = h
        set_cell_background(c, "173864")
        set_cell_margins(c, top=100, bottom=100, left=100, right=100)
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.runs[0].bold = True
        p.runs[0].font.name = 'Arial'
        p.runs[0].font.size = Pt(9.5)
        p.runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    for row_idx, r in enumerate(rows):
        bg = "F9FAFC" if row_idx % 2 == 1 else "FFFFFF"
        for col_idx, text in enumerate(r):
            c = act_table.rows[row_idx + 1].cells[col_idx]
            c.text = text
            set_cell_background(c, bg)
            set_cell_margins(c, top=80, bottom=80, left=80, right=80)
            p = c.paragraphs[0]
            p.runs[0].font.name = 'Arial'
            p.runs[0].font.size = Pt(9)

    doc.add_paragraph()

    # Section: Conclusión del periodo
    p_sec3 = doc.add_paragraph()
    r_sec3 = p_sec3.add_run("3. Estado Actual y Siguiente Fase")
    r_sec3.bold = True
    r_sec3.font.size = Pt(11)
    r_sec3.font.color.rgb = RGBColor(0x17, 0x38, 0x64)
    r_sec3.font.name = 'Arial'

    p_desc3 = doc.add_paragraph(
        "Con la culminación de la Fase 2, el proyecto cuenta con un diagnóstico situacional consolidado, delimitación metodológica aprobada y una matriz de evaluación de riesgos formalizada. La siguiente etapa (Fase 3: Aseguramiento de Red y Transporte) comprenderá el diseño topológico en GNS3, la configuración de zonas lógicas (VLANs), direccionamiento IP y políticas de filtrado mediante cortafuegos GNU/Linux."
    )
    p_desc3.runs[0].font.name = 'Arial'
    p_desc3.runs[0].font.size = Pt(10)

    out_file = '2026_II_SeguridadPerimetral_y_AseguramientoRedDatos/Bitacora_de_Cambios_Fase2.docx'
    doc.save(out_file)
    print(f"Bitacora guardada: {out_file}")

if __name__ == '__main__':
    create_bitacora_docx()
