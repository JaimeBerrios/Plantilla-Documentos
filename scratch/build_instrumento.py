import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

UNAB_BLUE = RGBColor(0x17, 0x38, 0x64)
TEXT_COLOR = RGBColor(0x22, 0x22, 0x22)

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=140, right=140):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def set_table_borders(table, color="173864", sz="4", val="single"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'  <w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:insideH w:val="none"/>'
        f'  <w:insideV w:val="none"/>'
        f'  <w:left w:val="none"/>'
        f'  <w:right w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def build_instrumento_docx():
    doc = docx.Document()

    # Set page margins (1 inch)
    section = doc.sections[0]
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)

    # Header block
    p_inst = doc.add_paragraph()
    p_inst.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_inst.paragraph_format.space_after = Pt(2)
    p_inst.paragraph_format.line_spacing = 1.15
    r_inst = p_inst.add_run("UNIVERSIDAD DR. ANDRÉS BELLO\n")
    r_inst.bold = True
    r_inst.font.name = 'Arial'
    r_inst.font.size = Pt(13)
    r_inst.font.color.rgb = UNAB_BLUE

    r_fac = p_inst.add_run("FACULTAD DE TECNOLOGÍA E INNOVACIÓN\nINGENIERÍA EN SISTEMAS Y COMPUTACIÓN\n")
    r_fac.bold = True
    r_fac.font.name = 'Arial'
    r_fac.font.size = Pt(10.5)
    r_fac.font.color.rgb = UNAB_BLUE

    r_asig = p_inst.add_run("Diseño de Sistemas de Seguridad de Redes de Datos — Ciclo 02-2026\n")
    r_asig.font.name = 'Arial'
    r_asig.font.size = Pt(10)
    r_asig.font.color.rgb = TEXT_COLOR

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(8)
    p_title.paragraph_format.space_after = Pt(12)
    r_title = p_title.add_run("INSTRUMENTO DE RECOLECCIÓN DE INFORMACIÓN TÉCNICA\nCUESTIONARIO DE DIAGNÓSTICO DE INFRAESTRUCTURA Y SEGURIDAD DE RED")
    r_title.bold = True
    r_title.font.name = 'Arial'
    r_title.font.size = Pt(12)
    r_title.font.color.rgb = UNAB_BLUE

    # Metadata table
    meta_table = doc.add_table(rows=5, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(meta_table, color="173864", sz="4", val="single")

    meta_items = [
        ("Unidad de análisis:", "Red de datos de la Sede de Enlaces con la Educación (Santa Rosa de Lima)."),
        ("Informantes clave:", "Personal Técnico y Encargada de la Sede Institucional."),
        ("Técnica metodológica:", "Entrevista estructurada y levantamiento técnico in situ (Hernández-Sampieri, 2023)."),
        ("Objetivo del instrumento:", "Obtener datos cuantitativos y descriptivos sobre topología, estado de obsolescencia del hardware, ausencia de segmentación lógica y carencia de controles perimetrales."),
        ("Equipo investigador:", "Genesis Saravia (Líder), Jaime Berrios, Jefferson Bonilla, Rosmeri González, Florentín Treminio.")
    ]

    for idx, (k, v) in enumerate(meta_items):
        c0, c1 = meta_table.rows[idx].cells
        c0.text = k
        c1.text = v
        set_cell_background(c0, "F0F4F8")
        set_cell_background(c1, "FFFFFF")
        set_cell_margins(c0, top=50, bottom=50, left=80, right=80)
        set_cell_margins(c1, top=50, bottom=50, left=80, right=80)
        c0.paragraphs[0].runs[0].bold = True
        c0.paragraphs[0].runs[0].font.name = 'Arial'
        c0.paragraphs[0].runs[0].font.size = Pt(9)
        c0.paragraphs[0].runs[0].font.color.rgb = UNAB_BLUE
        c1.paragraphs[0].runs[0].font.name = 'Arial'
        c1.paragraphs[0].runs[0].font.size = Pt(9)

    doc.add_paragraph()

    # Introduction text
    p_intro = doc.add_paragraph()
    p_intro.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_intro.paragraph_format.line_spacing = 1.3
    p_intro.paragraph_format.space_after = Pt(10)
    r_intro = p_intro.add_run(
        "Indicaciones: El presente instrumento tiene como finalidad recabar información operativa y técnica directa de la infraestructura de comunicaciones de la sede. Las preguntas están organizadas en tres dimensiones analíticas y sus resultados constituyen la base fáctica para el diagnóstico situacional, la matriz de riesgos y la formulación de la propuesta tecnológica."
    )
    r_intro.font.name = 'Arial'
    r_intro.font.size = Pt(10)
    r_intro.italic = True

    # Questions and findings data
    sections = [
        ("DIMENSIÓN 1: INFRAESTRUCTURA FÍSICA Y EQUIPAMIENTO CENTRAL", [
            (
                "Ítem 1.1",
                "¿Cuántas estaciones de trabajo operan actualmente en la sede y bajo qué modalidades se distribuye la conectividad hacia los puestos del personal?",
                "En la sede operan activamente aproximadamente 10 equipos de cómputo para uso de oficina, asignados uno a cada técnico. La conectividad se distribuye de manera heterogénea: algunos equipos se enlazan mediante cable estructurado hacia el switch o routers intermedios, mientras que otros acceden exclusivamente mediante enlaces inalámbricos Wi-Fi.",
                "Confirma el tamaño de la unidad de análisis (10 nodos) y la dispersión en los medios de acceso físico y lógico."
            ),
            (
                "Ítem 1.2",
                "¿Cuál es el dispositivo principal que concentra el tráfico del edificio y cuál es su tiempo estimado de operación continua?",
                "La distribución central converge en un switch de red ubicado en la caja principal del edificio. Dicho equipo se encuentra instalado desde la construcción de las instalaciones en el año 2003 (aproximadamente 23 años de servicio). Debido a su obsolescencia, presenta inestabilidad y múltiples puertos físicos dañados que han dejado de transmitir señal.",
                "Evidencia obsolescencia crítica de hardware como activo principal, sustentando el riesgo extremo en la matriz."
            ),
            (
                "Ítem 1.3",
                "¿Cómo se administran las conexiones salientes del módem del proveedor de servicios de Internet (Claro)?",
                "El módem de Claro cuenta con únicamente tres salidas físicas activas: una destinada al sistema de cámaras de seguridad, otra conectada a un router de distribución y la tercera conectada directamente a la computadora principal de la encargada de la sede. Ante la saturación, recientemente fue necesario incorporar un nuevo router adicional para abastecer a los demás puestos.",
                "Demuestra insuficiencia de puertos y proliferación de routers en cascada sin gestión unificada ni segmentación."
            )
        ]),
        ("DIMENSIÓN 2: TOPOLOGÍA Y SEGMENTACIÓN LÓGICA DE RED", [
            (
                "Ítem 2.1",
                "¿Qué topología física rige la comunicación interna y existe actualmente alguna división lógica (VLAN) para separar los flujos de tráfico?",
                "La red mantiene históricamente una topología en estrella concentrada en el switch. No obstante, a nivel lógico la arquitectura es totalmente plana: no existe segmentación por VLANs ni delimitación de subredes. Todos los equipos comparten un único dominio de difusión (broadcast) en el mismo rango de direccionamiento.",
                "Justifica metodológicamente la necesidad de diseñar un modelo de segmentación lógica por roles operativos."
            ),
            (
                "Ítem 2.2",
                "¿Cuál es el tratamiento y canal de red utilizado cuando ingresan laptops externas de estudiantes o docentes para soporte técnico?",
                "Los equipos que ingresan para mantenimiento o diagnóstico preventivo y correctivo se conectan a los mismos puntos de red o redes Wi-Fi compartidas con las estaciones del personal administrativo. No se cuenta con una zona desmilitarizada (DMZ) ni con una subred aislada para equipos en cuarentena.",
                "Valida la existencia de vectores de amenaza de movimiento lateral e infección por malware hacia registros sensibles."
            )
        ]),
        ("DIMENSIÓN 3: SEGURIDAD PERIMETRAL Y CONTROL DE ACCESO", [
            (
                "Ítem 3.1",
                "¿Existe actualmente algún servidor central o proxy activo encargado de la auditoría y el filtrado del tráfico hacia Internet?",
                "Actualmente no se dispone de ningún servidor de control. En años anteriores la sede disponía de un servidor central y un proxy físico dedicado al filtrado y administración del acceso a Internet. Sin embargo, a raíz de modificaciones en la infraestructura edilicia, dichos servidores fueron retirados, quedando únicamente el switch y el cableado estructurado.",
                "Fundamenta la propuesta tecnológica de reincorporar un servidor proxy SQUID y cortafuegos perimetral GNU/Linux."
            ),
            (
                "Ítem 3.2",
                "¿Se aplican listas de control de acceso (ACL), políticas de puertos o inspección criptográfica (HTTPS/TLS)?",
                "No se cuenta con listas de control de acceso configuradas ni políticas formales para restringir servicios o puertos vulnerables. El tráfico saliente y transversal no atraviesa mecanismos de inspección profunda ni políticas forzadas de certificados digitales o filtrado de contenido.",
                "Establece los requerimientos de diseño para las reglas de filtrado perimetral y políticas de seguridad informática (PSI)."
            )
        ])
    ]

    for dim_title, items in sections:
        p_sec = doc.add_paragraph()
        p_sec.paragraph_format.space_before = Pt(12)
        p_sec.paragraph_format.space_after = Pt(6)
        r_sec = p_sec.add_run(dim_title)
        r_sec.bold = True
        r_sec.font.name = 'Arial'
        r_sec.font.size = Pt(11)
        r_sec.font.color.rgb = UNAB_BLUE

        for item_code, question, finding, impact in items:
            t = doc.add_table(rows=4, cols=2)
            t.alignment = WD_TABLE_ALIGNMENT.CENTER
            set_table_borders(t, color="173864", sz="4", val="single")

            rows_info = [
                ("Identificador / Pregunta:", f"{item_code} — {question}"),
                ("Hallazgo técnico registrado:", finding),
                ("Impacto en la investigación:", impact),
                ("Estado / Veredicto técnico:", "Condición observada y verificada in situ")
            ]

            for r_i, (label, text) in enumerate(rows_info):
                c0, c1 = t.rows[r_i].cells
                c0.text = label
                c1.text = text
                set_cell_background(c0, "F0F4F8" if r_i % 2 == 0 else "F9FAFC")
                set_cell_background(c1, "FFFFFF")
                set_cell_margins(c0, top=50, bottom=50, left=70, right=70)
                set_cell_margins(c1, top=50, bottom=50, left=70, right=70)

                c0.paragraphs[0].runs[0].bold = True
                c0.paragraphs[0].runs[0].font.name = 'Arial'
                c0.paragraphs[0].runs[0].font.size = Pt(8.5)
                c0.paragraphs[0].runs[0].font.color.rgb = UNAB_BLUE

                c1.paragraphs[0].runs[0].font.name = 'Arial'
                c1.paragraphs[0].runs[0].font.size = Pt(8.5)
                if r_i == 3:
                    c1.paragraphs[0].runs[0].bold = True
                    c1.paragraphs[0].runs[0].font.color.rgb = RGBColor(0x00, 0x7A, 0x33)

            doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # Conformance and signoff
    p_concl = doc.add_paragraph()
    p_concl.paragraph_format.space_before = Pt(12)
    p_concl.paragraph_format.space_after = Pt(12)
    p_concl.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r_c = p_concl.add_run(
        "Certificación metodológica: Los hallazgos recabados a través del presente instrumento reflejan fielmente el estado tecnológico y de seguridad perimetral de la Sede de Enlaces con la Educación al momento del levantamiento (periodo 2026), sirviendo de insumo directo para el diseño de la arquitectura simulada en GNS3."
    )
    r_c.font.name = 'Arial'
    r_c.font.size = Pt(9.5)
    r_c.italic = True

    p_sig = doc.add_paragraph()
    p_sig.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_sig = p_sig.add_run("San Miguel, El Salvador — Equipo de Investigación de Cátedra UNAB")
    r_sig.bold = True
    r_sig.font.name = 'Arial'
    r_sig.font.size = Pt(9.5)
    r_sig.font.color.rgb = UNAB_BLUE

    out_docx = "2026_II_SeguridadPerimetral_y_AseguramientoRedDatos/Anexos/Instrumento_Cuestionario_Diagnostico_Tecnico.docx"
    doc.save(out_docx)
    print(f"Cuestionario Word guardado exitosamente: {out_docx}")

if __name__ == '__main__':
    build_instrumento_docx()
